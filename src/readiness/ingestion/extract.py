"""Turn bytes into text. Pure functions; no network. Unknown types are skipped, never guessed."""

from __future__ import annotations

import io
import json
from pathlib import PurePosixPath

TEXT_EXT = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".tsv",
    ".log",
    ".cfg",
    ".conf",
    ".ini",
    ".tf",
    ".hcl",
}
STRUCTURED_EXT = {".json", ".yaml", ".yml", ".toml"}


def extract_text(name: str, data: bytes, max_chars: int = 400_000) -> tuple[str, bool, str | None]:
    """Returns (text, truncated, error)."""
    ext = PurePosixPath(name.lower()).suffix
    try:
        if ext in TEXT_EXT or ext in STRUCTURED_EXT:
            text = data.decode("utf-8", errors="replace")
            if ext == ".json":
                try:
                    text = json.dumps(json.loads(text), indent=1)
                except ValueError:
                    pass
        elif ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext in {".xlsx", ".xlsm"}:
            text = _xlsx(data)
        else:
            return "", False, f"unsupported type {ext or '(none)'}"
    except Exception as exc:  # noqa: BLE001 - extraction failures are reported, not fatal
        return "", False, f"{type(exc).__name__}: {exc}"[:200]
    truncated = len(text) > max_chars
    return text[:max_chars], truncated, None


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
